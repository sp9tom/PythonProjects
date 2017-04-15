import random
import openpyxl
from openpyxl.styles import Font, Style, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

wb = openpyxl.Workbook()

sheet = wb.get_sheet_by_name('Sheet')
sheet.title = 'Pomiary'

# set styles
bold11Font = Font(size=11, bold=True)
bold11RedFont = Font(size=11, bold=True, color="cc0000")
align = Alignment(horizontal="center", vertical="center")
styleObj = Style(font=bold11Font, alignment=align)
styleObjRed = Style(font=bold11RedFont, alignment=align)

# append styles to top cells of columns
sheet['A1'].style = styleObj
sheet['B1'].style = styleObj

# set column names
sheet['A1'] = 'No.'
sheet['B1'] = 'Measure'

# set sheet properties
sheet.column_dimensions['A'].width = 6
sheet.column_dimensions['B'].width = 10

# insert 'No' values
for no in range(2, 12):
    sheet['A'+str(no)] = no-1

# insert random values to column 'Measure'
for i in range(2, 12):
    sheet['B'+str(i)] = random.randint(0, 10)

# append styles to sum cells
sheet['A13'].style = styleObjRed
sheet['B13'].style = styleObjRed

# insert Sum cell
sheet['A13'] = 'Sum:'
sheet['B13'] = '=SUM(B2:B11)'

# prepare data for Excel charts
values = Reference(sheet, 2, 1, 2, 11)
points = Reference(sheet, 1, 2, 1, 11)

# create Excel chart
chart1 = BarChart()
chart1.type = "col"
chart1.style = 12
chart1.shape = 1
chart1.title = "Random Bar Chart"
chart1.x_axis.title = 'Number'
chart1.y_axis.title = 'Value'

chart1.add_data(values, titles_from_data=True)
chart1.set_categories(points)

sheet.add_chart(chart1, "D3")

# create second Excel chart
chart2 = LineChart()
chart2.type = "col"
chart2.style = 3
chart2.shape = 1
chart2.title = "Random Line Chart"
chart2.x_axis.title = 'Number'
chart2.y_axis.title = 'Value'

chart2.add_data(values, titles_from_data=True)
chart2.set_categories(points)

sheet.add_chart(chart2, "L3")

# save Excel sheet
wb.save("Example.xlsx")

print "Excel Sheet saved on the disk"
